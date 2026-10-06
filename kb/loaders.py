"""Turn uploaded bytes into text sections.

Each loader returns a LoadedDocument whose sections carry location metadata
(e.g. PDF page numbers) that flows through to chunk citations.
"""

import csv
import io
import json
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Callable

if TYPE_CHECKING:
    from .ocr import PdfOcr

# pypdf logs a warning per font it can't fully decode, often hundreds per document.
logging.getLogger("pypdf").setLevel(logging.ERROR)


class UnsupportedFileType(ValueError):
    pass


@dataclass
class Section:
    text: str
    metadata: dict = field(default_factory=dict)


@dataclass
class LoadedDocument:
    file_type: str
    sections: list[Section]
    # Document-level metadata discovered from the file itself (e.g. title).
    metadata: dict = field(default_factory=dict)
    # Whether the text uses markdown-style headings (affects chunk separators).
    structured: bool = False


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def load_pdf(data: bytes, ocr: "PdfOcr | None" = None) -> LoadedDocument:
    from pypdf import PdfReader

    if ocr is None:
        from .ocr import get_pdf_ocr

        ocr = get_pdf_ocr()

    reader = PdfReader(io.BytesIO(data))
    sections = [
        Section(page.extract_text() or "", {"page": i})
        for i, page in enumerate(reader.pages, start=1)
    ]
    # Scanned pages have no text layer: read them with OCR instead.
    scanned = [i for i, s in enumerate(sections) if ocr.needs_ocr(s.text)]
    if scanned:
        for i, text in ocr.read_pages(data, scanned).items():
            sections[i].text = text
            sections[i].metadata["ocr"] = True
    meta = {}
    if reader.metadata and reader.metadata.title:
        meta["title"] = str(reader.metadata.title)
    return LoadedDocument("pdf", sections, meta)


def load_docx(data: bytes) -> LoadedDocument:
    import docx
    from docx.table import Table

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
            lines.append("")
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        if style.startswith("Heading"):
            level = style.removeprefix("Heading").strip()
            hashes = "#" * int(level) if level.isdigit() else "#"
            lines.append(f"\n{hashes} {text}")
        elif style == "Title":
            lines.append(f"\n# {text}")
        else:
            lines.append(text)
        lines.append("")
    meta = {}
    if document.core_properties.title:
        meta["title"] = document.core_properties.title
    return LoadedDocument("docx", [Section("\n".join(lines))], meta, structured=True)


def load_html(data: bytes) -> LoadedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None
    for tag in soup(["head", "script", "style", "noscript", "nav", "footer", "header", "form"]):
        tag.decompose()
    for level in range(1, 7):
        for heading in soup.find_all(f"h{level}"):
            heading.string = f"\n\n{'#' * level} {heading.get_text(' ', strip=True)}\n\n"
    raw = soup.get_text("\n")
    lines = [line.strip() for line in raw.splitlines()]
    text = "\n".join(lines)
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    meta = {"title": title} if title else {}
    return LoadedDocument("html", [Section(text)], meta, structured=True)


def load_markdown(data: bytes) -> LoadedDocument:
    return LoadedDocument("markdown", [Section(_decode(data))], structured=True)


def load_text(data: bytes) -> LoadedDocument:
    return LoadedDocument("text", [Section(_decode(data))])


def load_csv(data: bytes) -> LoadedDocument:
    reader = csv.DictReader(io.StringIO(_decode(data)))
    rows = []
    for row in reader:
        pairs = [f"{k}: {v}" for k, v in row.items() if k and v not in (None, "")]
        if pairs:
            rows.append(" | ".join(pairs))
    return LoadedDocument("csv", [Section("\n".join(rows))])


def load_json(data: bytes) -> LoadedDocument:
    parsed = json.loads(_decode(data))
    return LoadedDocument("json", [Section(json.dumps(parsed, indent=2, ensure_ascii=False))])


LOADERS: dict[str, Callable[[bytes], LoadedDocument]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".html": load_html,
    ".htm": load_html,
    ".md": load_markdown,
    ".markdown": load_markdown,
    ".txt": load_text,
    ".csv": load_csv,
    ".json": load_json,
}

SUPPORTED_EXTENSIONS = sorted(LOADERS)


def load_document(filename: str, data: bytes) -> LoadedDocument:
    ext = Path(filename).suffix.lower()
    loader = LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext or filename}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    return loader(data)

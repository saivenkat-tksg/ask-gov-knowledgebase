"""Turn uploaded bytes into text sections.

Each loader returns a LoadedDocument whose sections carry location metadata
(e.g. PDF page numbers) that flows through to chunk citations.
"""

import csv
import io
import json
import logging
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING

from kb.ingestion.tables import markdown_table

if TYPE_CHECKING:
    from kb.ingestion.ocr import PdfOcr

# pypdf logs a warning per font it can't fully decode, often hundreds per document.
logging.getLogger("pypdf").setLevel(logging.ERROR)
log = logging.getLogger(__name__)


class UnsupportedFileType(ValueError):
    pass


@dataclass
class Section:
    text: str
    metadata: dict = field(default_factory=dict)


@dataclass
class LoadedDocument:
    file_type: str
    sections: list[Section]
    # Document-level metadata discovered from the file itself (e.g. title).
    metadata: dict = field(default_factory=dict)
    # Whether the text uses markdown-style headings (affects chunk separators).
    structured: bool = False
    # PDF headings found from fonts and marked as markdown headings (structure=True only).
    headings_marked: int = 0
    tables: int = 0  # tables written as markdown tables (structure=True, tables=True only)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def load_pdf(data: bytes, ocr: "PdfOcr | None" = None, structure: bool = False, tables: bool = True) -> LoadedDocument:
    """`structure`: rebuild ruled tables as markdown tables (when `tables`), then find headings from
    fonts and mark them as markdown headings, for structure-aware chunking."""
    from pypdf import PdfReader

    if ocr is None:
        from kb.ingestion.ocr import get_pdf_ocr

        ocr = get_pdf_ocr()

    reader = PdfReader(io.BytesIO(data))
    sections = [
        Section(page.extract_text() or "", {"page": i})
        for i, page in enumerate(reader.pages, start=1)
    ]
    # Scanned pages have no text layer: read them with OCR instead.
    scanned = [i for i, s in enumerate(sections) if ocr.needs_ocr(s.text)]
    if scanned:
        for i, text in ocr.read_pages(data, scanned).items():
            sections[i].text = text
            sections[i].metadata["ocr"] = True
    meta = {}
    if reader.metadata and reader.metadata.title:
        meta["title"] = str(reader.metadata.title)
    marked = found_tables = 0
    if structure and tables:
        from kb.ingestion.pdf_tables import tables_by_page

        try:
            for i, (text, count) in tables_by_page(data, skip=set(scanned)).items():
                sections[i].text = text
                sections[i].metadata["tables"] = count
                found_tables += count
        except Exception as exc:
            log.warning("PDF table extraction failed, keeping plain text: %s", exc)
    if structure:
        from kb.ingestion.pdf_headings import detect_headings as find_headings
        from kb.ingestion.pdf_headings import mark_headings, read_lines, strip_running_lines

        try:
            texts, marked = mark_headings([s.text for s in sections], find_headings(read_lines(data)))
            texts = strip_running_lines(texts)
        except Exception as exc:  # a PDF pdfium can't read: keep the text, chunk without structure
            log.warning("PDF heading detection failed, chunking without structure: %s", exc)
        else:
            for section, text in zip(sections, texts):
                section.text = text
    return LoadedDocument("pdf", sections, meta, structured=marked > 0, headings_marked=marked, tables=found_tables)


def load_docx(data: bytes, tables: bool = False) -> LoadedDocument:
    """`tables`: write tables as markdown tables (first row as header) instead of `a | b` lines."""
    import docx
    from docx.table import Table

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    count = 0
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            rows = []
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if tables:  # a merged cell repeats its text in every column it spans
                    cells = [c if k == 0 or c != cells[k - 1] else "" for k, c in enumerate(cells)]
                if any(cells):
                    rows.append(cells)
            if tables and len(rows) >= 2:
                lines.append(markdown_table(rows[0], rows[1:]))
                count += 1
            else:
                lines.extend(" | ".join(cells) for cells in rows)
            lines.append("")
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        if style.startswith("Heading"):
            level = style.removeprefix("Heading").strip()
            hashes = "#" * int(level) if level.isdigit() else "#"
            lines.append(f"\n{hashes} {text}")
        elif style == "Title":
            lines.append(f"\n# {text}")
        else:
            lines.append(text)
        lines.append("")
    meta = {}
    if document.core_properties.title:
        meta["title"] = document.core_properties.title
    return LoadedDocument("docx", [Section("\n".join(lines))], meta, structured=True, tables=count)


def load_html(data: bytes, tables: bool = False) -> LoadedDocument:
    """`tables`: write <table>s as markdown tables instead of one line per cell."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None
    for tag in soup(["head", "script", "style", "noscript", "nav", "footer", "header", "form"]):
        tag.decompose()
    count = 0
    if tables:
        for table in soup.find_all("table"):
            if table.find_parent("table"):
                continue  # nested tables are flattened into their outer cell
            rows = [[c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])] for tr in table.find_all("tr")]
            rows = [r for r in rows if any(r)]
            if len(rows) >= 2:
                caption = table.find("caption")
                md = markdown_table(rows[0], rows[1:], caption.get_text(" ", strip=True) if caption else None)
                table.replace_with(f"\n\n{md}\n\n")
                count += 1
    for level in range(1, 7):
        for heading in soup.find_all(f"h{level}"):
            heading.string = f"\n\n{'#' * level} {heading.get_text(' ', strip=True)}\n\n"
    raw = soup.get_text("\n")
    lines = [line.strip() for line in raw.splitlines()]
    text = "\n".join(lines)
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    meta = {"title": title} if title else {}
    return LoadedDocument("html", [Section(text)], meta, structured=True, tables=count)


def load_markdown(data: bytes) -> LoadedDocument:
    return LoadedDocument("markdown", [Section(_decode(data))], structured=True)


def load_text(data: bytes) -> LoadedDocument:
    return LoadedDocument("text", [Section(_decode(data))])


def load_csv(data: bytes) -> LoadedDocument:
    reader = csv.DictReader(io.StringIO(_decode(data)))
    rows = []
    for row in reader:
        pairs = [f"{k}: {v}" for k, v in row.items() if k and v not in (None, "")]
        if pairs:
            rows.append(" | ".join(pairs))
    return LoadedDocument("csv", [Section("\n".join(rows))])


def load_json(data: bytes) -> LoadedDocument:
    parsed = json.loads(_decode(data))
    return LoadedDocument("json", [Section(json.dumps(parsed, indent=2, ensure_ascii=False))])


def document_text(sections: list[Section]) -> tuple[str, list[list[int]]]:
    """The loader's sections (PDF pages) joined as one text, with a [offset, page] map."""
    parts, page_map, offset = [], [], 0
    for section in sections:
        page = section.metadata.get("page")
        if page is not None:
            page_map.append([offset, page])
        parts.append(section.text)
        offset += len(section.text) + 2
    return "\n\n".join(parts), page_map


LOADERS: dict[str, Callable[[bytes], LoadedDocument]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".html": load_html,
    ".htm": load_html,
    ".md": load_markdown,
    ".markdown": load_markdown,
    ".txt": load_text,
    ".csv": load_csv,
    ".json": load_json,
}

SUPPORTED_EXTENSIONS = sorted(LOADERS)


def load_document(filename: str, data: bytes, structure: bool = False, tables: bool = True) -> LoadedDocument:
    """`structure` prepares text for structure-aware chunking: PDF headings found from fonts become
    markdown headings, and (with `tables`) PDF, DOCX and HTML tables become markdown tables.
    With structure=False every loader returns exactly what it always has."""
    ext = Path(filename).suffix.lower()
    loader = LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext or filename}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    if ext == ".pdf":
        return load_pdf(data, structure=structure, tables=tables)
    if ext == ".docx":
        return load_docx(data, tables=structure and tables)
    if ext in (".html", ".htm"):
        return load_html(data, tables=structure and tables)
    return loader(data)

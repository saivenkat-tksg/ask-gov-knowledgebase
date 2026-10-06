import io
import json

import pytest

from kb.config import Settings
from kb.loaders import UnsupportedFileType, load_document, load_pdf
from kb.ocr import PdfOcr


def test_text_and_markdown():
    assert load_document("a.txt", b"hello").sections[0].text == "hello"
    md = load_document("a.md", "# Title\nbody".encode())
    assert md.structured and md.file_type == "markdown"


def test_csv_rows_become_lines():
    doc = load_document("a.csv", b"name,age\nAnn,30\nBob,\n")
    assert doc.sections[0].text == "name: Ann | age: 30\nname: Bob"


def test_json():
    doc = load_document("a.json", json.dumps({"k": [1, 2]}).encode())
    assert '"k"' in doc.sections[0].text


def test_html_strips_scripts_and_marks_headings():
    html = b"<html><head><title>T</title><script>x=1</script></head><body><h2>Sec</h2><p>Body</p></body></html>"
    doc = load_document("a.html", html)
    assert doc.metadata["title"] == "T"
    assert "## Sec" in doc.sections[0].text and "x=1" not in doc.sections[0].text


def test_docx_headings_and_tables():
    import docx

    d = docx.Document()
    d.add_heading("Intro", level=1)
    d.add_paragraph("Some text.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "a", "b"
    buf = io.BytesIO()
    d.save(buf)
    text = load_document("a.docx", buf.getvalue()).sections[0].text
    assert "# Intro" in text and "Some text." in text and "a | b" in text


def test_pdf_pages():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(100, 100)
    w.add_blank_page(100, 100)
    buf = io.BytesIO()
    w.write(buf)
    doc = load_pdf(buf.getvalue(), ocr=PdfOcr(Settings(ocr_mode="off")))
    assert [s.metadata["page"] for s in doc.sections] == [1, 2]


def test_unsupported():
    with pytest.raises(UnsupportedFileType):
        load_document("a.exe", b"")

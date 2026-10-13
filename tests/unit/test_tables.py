import io

from kb.ingestion.loaders import load_document
from kb.ingestion.pdf_headings import Heading, mark_headings
from kb.ingestion.pdf_tables import build_table
from kb.ingestion.structure import StructureChunker
from kb.ingestion.tables import is_row, markdown_table, split_blocks, split_table


def w(text: str, x0: float, top: float, bold: bool = False) -> dict:
    return {"text": text, "x0": x0, "x1": x0 + 6 * len(text), "top": top, "bottom": top + 10,
            "fontname": "Poppins-Bold" if bold else "Poppins-Regular"}


def test_markdown_table_format():
    md = markdown_table(["Step", "Action"], [["1.", "Pay | sign"], ["2."]], title="Fees")
    assert md.splitlines() == [
        "Table: Fees",
        "| Step | Action |",
        "| --- | --- |",
        "| 1. | Pay / sign |",  # a pipe inside a cell can't break the row
        "| 2. |  |",  # short rows are padded
    ]


def test_split_blocks_finds_tables_with_titles():
    text = "Intro line.\nTable: Fees\n| A | B |\n| --- | --- |\n| 1 | 2 |\nAfter."
    kinds = [(b.kind, text[b.start:b.end].strip()) for b in split_blocks(text)]
    assert kinds[0] == ("text", "Intro line.")
    assert kinds[1][0] == "table" and kinds[1][1].startswith("Table: Fees") and kinds[1][1].endswith("| 1 | 2 |")
    assert kinds[2] == ("text", "After.")


def test_split_table_between_rows_repeating_header():
    rows = [[f"{i}.", f"step number {i} " + "word " * 8] for i in range(1, 21)]
    md = markdown_table(["Step", "Action"], rows, title="Process")
    pieces = split_table(md, 200, len)
    assert len(pieces) > 2
    seen = []
    for text, _, _ in pieces:
        lines = text.splitlines()
        assert len(text) <= 200
        assert lines[:3] == ["Table: Process", "| Step | Action |", "| --- | --- |"]
        assert all(is_row(line) for line in lines[1:])
        seen += lines[3:]
    assert seen == md.splitlines()[3:]  # every row exactly once, in order
    assert split_table(md, 10_000, len) == [(md, 0, len(md))]  # a table that fits is untouched


def test_build_table_from_words():
    words = [
        w("ID", 80, 10, True), w("Card", 100, 10, True), w("Creation", 130, 10, True),  # title
        w("Step", 80, 30, True), w("Action", 140, 30, True), w("Completed", 330, 30, True), w("Notes", 450, 30, True),
        w("by", 330, 42, True), w("(Role)", 348, 42, True),  # header continues on a second line
        w("1.", 80, 60), w("Form", 140, 60), w("is", 172, 60), w("emailed", 188, 60), w("Clerk", 330, 60),
        w("to", 140, 72), w("supervisor", 156, 72),  # wrapped line joins the row above
        w("2.", 80, 90), w("Photo", 140, 90), w("clipped", 178, 90), w("Supervisor", 330, 90),
        w("Uses", 450, 90), w("CorelDRAW", 480, 90),
    ]
    t = build_table(words, (70, 0, 560, 100))
    assert t.title == "ID Card Creation"
    assert t.header == ["Step", "Action", "Completed by (Role)", "Notes"]
    assert t.rows == [["1.", "Form is emailed to supervisor", "Clerk", ""],
                      ["2.", "Photo clipped", "Supervisor", "Uses CorelDRAW"]]
    assert not t.continued


def test_build_table_continuation_and_step_column():
    words = [w("4.", 60, 10), w("Officer", 100, 10), w("escorts", 148, 10), w("Security", 330, 10),
             w("5.", 60, 30), w("Evidence", 100, 30), w("submitted", 154, 30), w("Police", 330, 30)]
    t = build_table(words, (50, 0, 560, 50))
    assert t.continued  # starts with a step number: the header is on the previous page
    assert t.rows == [["4.", "Officer escorts", "Security"], ["5.", "Evidence submitted", "Police"]]


def test_one_column_box_is_not_a_table():
    words = [w("A", 80, 10), w("bordered", 92, 10), w("note", 146, 10), w("continues", 80, 22), w("here", 140, 22)]
    assert build_table(words, (70, 0, 300, 40)) is None


def test_chunker_splits_large_table_by_rows():
    rows = [[f"{i}.", "action text " * 6, "Clerk"] for i in range(1, 30)]
    text = "# Guide\n\n## Process\n\nIntro text.\n\n" + markdown_table(["#", "Action", "Role"], rows, "Steps") + "\n\n## Other\n\nx\n\n## More\n\ny"
    chunks = StructureChunker(chunk_size=400, chunk_overlap=0, min_tokens=10).split(text)
    table_chunks = [c for c in chunks if "| --- |" in c.text]
    assert len(table_chunks) > 2
    for c in table_chunks:
        assert c.text.startswith("[Guide > Process]")
        assert "| # | Action | Role |" in c.text
        assert all(line.endswith("|") for line in c.text.splitlines() if line.startswith("|"))
        assert len(c.text) <= 400


def test_heading_in_title_row_splits_table():
    page = ("Table: 1. TINT Waiver\n| # | Capability | Benefits |\n| --- | --- | --- |\n"
            "| 1. | Online form | Faster |\n| 2. | Firearms Registration |  |\n| 1. | Online application | Easier |\n")
    out, found = mark_headings([page], [Heading(0, "1. TINT Waiver", 3), Heading(0, "2. Firearms Registration", 3),
                                        Heading(0, "Faster", 4)])
    text = out[0]
    assert found == 2  # "Faster" is an ordinary cell: never breaks a row
    assert text.index("### 1. TINT Waiver") < text.index("Table: 1. TINT Waiver")  # heading above the table
    after = text.split("### 2. Firearms Registration", 1)[1].strip().splitlines()
    assert after[:3] == ["| # | Capability | Benefits |", "| --- | --- | --- |", "| 1. | Online application | Easier |"]


def test_docx_and_html_tables_become_markdown_with_structure():
    html = (b"<html><body><h1>Fees</h1><table><caption>Certificate fees</caption>"
            b"<tr><th>Service</th><th>Fee</th></tr><tr><td>Birth</td><td>GY$1,000</td></tr></table></body></html>")
    doc = load_document("fees.html", html, structure=True)
    assert "Table: Certificate fees\n| Service | Fee |\n| --- | --- |\n| Birth | GY$1,000 |" in doc.sections[0].text
    assert doc.tables == 1
    plain = load_document("fees.html", html)  # recursive strategy: loader output unchanged
    assert "| Service |" not in plain.sections[0].text and plain.tables == 0

    import docx

    d = docx.Document()
    table = d.add_table(rows=2, cols=2)
    for r, values in enumerate([("Service", "Fee"), ("Death", "GY$500")]):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    buf = io.BytesIO()
    d.save(buf)
    loaded = load_document("fees.docx", buf.getvalue(), structure=True)
    assert "| Service | Fee |\n| --- | --- |\n| Death | GY$500 |" in loaded.sections[0].text
